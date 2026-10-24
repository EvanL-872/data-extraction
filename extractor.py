"""In-memory extraction of purchase-request workbooks in the observed A:I layout."""

from __future__ import annotations

import io
import math
import re
from collections.abc import Iterable, Mapping
from typing import Any

import pandas as pd
from openpyxl import load_workbook


ITEM_COLUMNS = [
    "file", "sheet", "customer", "sales_rep", "block_id", "supplier", "sn", "sku",
    "item_name", "description", "item_type", "brand_category", "segment", "section",
    "start_row", "end_row", "cost_currency", "sell_currency", "cost_fx_to_sgd",
    "sell_fx_to_sgd", "cost_fx_source", "sell_fx_source", "total_cost", "total_sell", "total_cost_sgd", "total_sell_sgd",
    "missing_cost_rows", "missing_sell_rows",
]
SUPPLIER_COLUMNS = [
    "file", "sheet", "customer", "sales_rep", "block_id", "supplier", "cost_currency",
    "sell_currency", "cost_fx_to_sgd", "total_cost", "total_sell", "total_cost_sgd",
    "total_sell_sgd", "cost_fx_source", "sell_fx_source", "declared_cost", "cost_variance", "missing_cost_rows",
    "missing_sell_rows",
]
LINE_COLUMNS = [
    "file", "sheet", "block_id", "supplier", "item_start_row", "excel_row", "sn", "sku",
    "description", "line_item_type", "brand_category", "quantity", "cost_raw", "sell_raw",
    "cost", "sell", "cost_status", "sell_status", "cost_currency", "sell_currency",
]
ISSUE_COLUMNS = ["file", "sheet", "excel_row", "severity", "issue"]
ADJUSTMENT_COLUMNS = ["file", "sheet", "excel_row", "description", "sell", "cost", "sell_currency"]
RECONCILIATION_COLUMNS = ["file", "sheet", "side", "recorded_sgd", "pr_summary", "variance", "missing_item_amounts"]
USD_TO_SGD_BASELINE = 1.33


class TemplateError(ValueError):
    """Raised when a selected worksheet is not a supported PR template."""


def clean_text(value: Any) -> str:
    return "" if value is None else " ".join(str(value).split())


def as_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        number = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def find_currency(value: Any) -> str | None:
    match = re.search(r"\b(SGD|USD|EUR|GBP|AUD|MYR|JPY|CNY|HKD)\b", clean_text(value).upper())
    return match.group(1) if match else None


def sum_known(values: Iterable[float | None]) -> float | None:
    numbers = [value for value in values if value is not None and pd.notna(value)]
    return sum(numbers) if numbers else None


def product_type(sku: Any, description: Any) -> str:
    """Classify at a business level; brand classification remains a separate field."""
    label = f"{clean_text(sku)} {clean_text(description)}"
    if re.search(r"license|licence|subscription|\bLIC-", label, re.I):
        return "License / Subscription"
    if re.search(r"warranty|maintenance|support service|\bRMA\b|\bSVC\b", label, re.I):
        return "Support / Warranty"
    if re.search(r"discount|rebate|credit note", label, re.I):
        return "Discount"
    if re.search(
        r"professional\s+services?|\bPS-|supply\s*(?:&|and)\s*install|provide services?|"
        r"installation|reinstall|configurations?|cabling works|network cabling|project management|"
        r"supply labour|mounting of|commissioning|programming|integration service",
        label,
        re.I,
    ):
        return "Professional Service"
    return "Hardware / Other"


BRAND_PATTERNS: tuple[tuple[str, str], ...] = (
    ("Cisco", r"\b(?:cisco|meraki)\b|\b(?:MR|MS|MX|MV|MT|CW)\d{2,}|\bcatalyst\b"),
    ("Aruba", r"\baruba\b|\bhpe\s+aruba\b|\binstant on\b"),
    ("Ruckus", r"\bruckus\b|\bICX[- ]?\d+"),
    ("Logitech", r"\b(?:logitech|logi)\b|\brally\b|\bmeetup\b|\bbrio\b|\btap(?:\s+ip)?\b"),
    ("Poly", r"\bpoly(?:com)?\b|\bstudio\s*x\b"),
    ("Yealink", r"\byealink\b"),
    ("Aver", r"\baver(?:\s+information)?\b|\bvc5\d{2}\b"),
    ("Samsung", r"\bsamsung\b"),
    ("APC", r"\bapc\b|\bschneider electric\b"),
    ("Ubiquiti", r"\bubiquiti\b|\bunifi\b"),
    ("Fortinet", r"\bfortinet\b|\bfortigate\b"),
    ("Palo Alto", r"\bpalo alto\b|\bpan-?\d{3,}"),
    ("Lenovo", r"\blenovo\b|\bthinkpad\b|\bthinkcentre\b"),
    ("Dell", r"\bdell\b|\bpoweredge\b"),
    ("HP", r"\bhewlett[ -]packard\b|\bHP\s+(?:pro|elite|zbook|laserjet)"),
    ("Microsoft", r"\bmicrosoft\b|\bwindows\s+(?:server|11|10)\b|\boffice 365\b"),
    ("Zoom", r"\bzoom rooms?\b"),
)


def brand_category(sku: Any, description: Any, item_kind: str | None = None) -> str:
    if item_kind == "Professional Service":
        return "Professional Service"
    label = f"{clean_text(sku)} {clean_text(description)}"
    for brand, pattern in BRAND_PATTERNS:
        if re.search(pattern, label, re.I):
            return brand
    if item_kind == "License / Subscription":
        return "License / Subscription"
    if item_kind == "Support / Warranty":
        return "Support / Warranty"
    if item_kind == "Discount":
        return "Discount"
    return "Unclassified"


def _is_pr_header(values: tuple[Any, ...] | list[Any]) -> bool:
    if len(values) < 5 or clean_text(values[0]).lower() != "s/n":
        return False
    return [clean_text(value).upper() for value in values[1:5]] == ["SKU", "DESCRIPTION", "QTY", "UOM"]


def _header_row(worksheet: Any) -> int | None:
    for row in range(1, min(worksheet.max_row, 250) + 1):
        values = [worksheet.cell(row, column).value for column in range(1, 6)]
        if _is_pr_header(values):
            return row
    return None


def discover_pr_sheets(content: bytes) -> list[str]:
    """Return only worksheets that expose the supported PR item header."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        matches = []
        for worksheet in workbook.worksheets:
            for values in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 250), max_col=5, values_only=True):
                if _is_pr_header(values):
                    matches.append(worksheet.title)
                    break
        return matches
    finally:
        workbook.close()


def read_pr_metadata(content: bytes, sheet_name: str) -> dict[str, str]:
    """Read display-only metadata for the sheet selector without persisting the upload."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        header = _header_row(worksheet)
        if header is None:
            return {}
        metadata: dict[str, str] = {}
        for row in range(1, header):
            label = clean_text(worksheet.cell(row, 1).value).rstrip(":").lower()
            value = clean_text(worksheet.cell(row, 2).value)
            if label and value:
                metadata[label] = value
        return {
            "customer": metadata.get("company name", ""),
            "sales_rep": metadata.get("sales rep", ""),
            "sell_currency": find_currency(metadata.get("currency")) or "",
        }
    finally:
        workbook.close()


def _empty_frame(columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=columns)


def extract_pr(
    content: bytes,
    filename: str,
    sheet_name: str,
    sell_currency_override: str | None = None,
    sell_fx_to_sgd: float | None = None,
    category_overrides: Mapping[int, str] | None = None,
) -> dict[str, pd.DataFrame]:
    """Extract one worksheet without writing the source workbook to local storage.

    The observed template places Supplier Details after the items belonging to that
    supplier. A numbered S/N begins an item; unnumbered continuation rows stay with
    it. Price fields are read as cached Excel values, so formulas without a cached
    value are surfaced in the review sheet instead of being silently replaced.
    """
    cached = load_workbook(io.BytesIO(content), data_only=True)
    formulas = load_workbook(io.BytesIO(content), data_only=False)
    try:
        if sheet_name not in cached.sheetnames:
            raise TemplateError(f"Worksheet '{sheet_name}' does not exist in {filename}.")
        worksheet, formula_sheet = cached[sheet_name], formulas[sheet_name]
        header = _header_row(worksheet)
        if header is None:
            raise TemplateError("No supported S/N, SKU, Description, Qty, UOM header was found.")

        rows = {
            row: [worksheet.cell(row, column).value for column in range(1, 10)]
            for row in range(1, worksheet.max_row + 1)
        }
        metadata = {
            clean_text(values[0]).rstrip(":").lower(): values[1]
            for row, values in rows.items()
            if row < header and clean_text(values[0])
        }
        customer = clean_text(metadata.get("company name"))
        sales_rep = clean_text(metadata.get("sales rep"))
        sell_currency = (sell_currency_override or find_currency(metadata.get("currency")) or "").upper() or None
        metadata_sell_rate = as_number(metadata.get("exchange rate"))
        if metadata_sell_rate is not None and metadata_sell_rate <= 0:
            metadata_sell_rate = None
        stop = next(
            (row for row, values in rows.items() if clean_text(values[0]).upper() == "OTHER COST"),
            worksheet.max_row + 1,
        )
        markers = [
            row for row, values in rows.items()
            if header < row < stop and clean_text(values[1]).lower() == "supplier details"
        ]
        if not markers:
            raise TemplateError("No trailing 'Supplier Details' blocks found on this worksheet.")

        issues: list[dict[str, Any]] = []
        items: list[dict[str, Any]] = []
        lines: list[dict[str, Any]] = []
        suppliers: list[dict[str, Any]] = []
        adjustments: list[dict[str, Any]] = []
        reconciliation: list[dict[str, Any]] = []

        def warn(row: int, message: str, severity: str = "Warning") -> None:
            issues.append({"file": filename, "sheet": sheet_name, "excel_row": row, "severity": severity, "issue": message})

        start = header + 1
        known_supplier_labels = {
            "company name", "company address", "contact person", "email", "payment terms",
            "currency", "exchange rate",
        }
        for marker in markers:
            fields: dict[str, Any] = {}
            declared: dict[str, float | None] = {}
            block_end = marker
            for row in range(marker + 1, min(marker + 20, worksheet.max_row + 1)):
                values = rows[row]
                label = clean_text(values[1]).rstrip(":").lower()
                if label in known_supplier_labels:
                    fields[label] = values[2]
                    block_end = row
                elif label.startswith("total supplier cost"):
                    declared[find_currency(label) or "native"] = as_number(values[2])
                    block_end = row
                elif not any(clean_text(value) for value in values[:3]):
                    continue
                else:
                    break

            supplier = clean_text(fields.get("company name")) or "UNKNOWN SUPPLIER"
            cost_currency = find_currency(fields.get("currency"))
            stated_cost_rate = as_number(fields.get("exchange rate"))
            if stated_cost_rate is not None and stated_cost_rate <= 0:
                stated_cost_rate = None
            if cost_currency == "SGD":
                cost_rate, cost_fx_source = 1.0, "SGD fixed rate"
            elif stated_cost_rate is not None:
                cost_rate, cost_fx_source = stated_cost_rate, "PR supplier exchange rate"
            elif cost_currency == "USD":
                cost_rate, cost_fx_source = USD_TO_SGD_BASELINE, "USD baseline 1.33 (no PR rate)"
                warn(marker, "No USD supplier exchange rate in PR; used SGD reference baseline 1.33.", "Info")
            else:
                cost_rate, cost_fx_source = None, "Unavailable"
            if not cost_currency or cost_rate is None:
                warn(marker, "Missing supplier currency or valid SGD exchange rate; converted cost is blank.")
            if not sell_currency:
                warn(marker, "Missing customer selling currency; set an override before relying on SGD revenue.")
            if sell_currency == "SGD":
                sell_rate, sell_fx_source = 1.0, "SGD fixed rate"
            elif sell_fx_to_sgd is not None and sell_fx_to_sgd > 0:
                sell_rate, sell_fx_source = sell_fx_to_sgd, "Caller-supplied sell exchange rate"
            elif metadata_sell_rate is not None:
                sell_rate, sell_fx_source = metadata_sell_rate, "PR selling exchange rate"
            elif sell_currency and sell_currency == cost_currency and cost_rate is not None:
                sell_rate, sell_fx_source = cost_rate, f"Same-currency supplier rate: {cost_fx_source}"
            elif sell_currency == "USD":
                sell_rate, sell_fx_source = USD_TO_SGD_BASELINE, "USD baseline 1.33 (no PR rate)"
                warn(marker, "No USD selling exchange rate in PR; used SGD reference baseline 1.33.", "Info")
            else:
                sell_rate, sell_fx_source = None, "Unavailable"
            if sell_currency and sell_rate is None:
                warn(marker, "Missing sell-side SGD rate; native sell prices are retained but SGD revenue is blank.")
            if sell_fx_to_sgd is not None and sell_fx_to_sgd <= 0:
                warn(marker, "Caller-supplied sell-side SGD rate was not positive and was ignored.")

            block_id = f"{filename}|{sheet_name}|supplier@{marker}"
            groups: list[dict[str, Any]] = []
            current: dict[str, Any] | None = None
            section = ""

            def new_group(row: int, values: list[Any], serial: str = "", grouped: bool = False) -> dict[str, Any]:
                group = {
                    "start": row,
                    "sn": serial,
                    "sku": clean_text(values[1]),
                    "name": clean_text(values[2]) or clean_text(values[1]),
                    "section": section,
                    "grouped": grouped,
                    "rows": [],
                }
                groups.append(group)
                return group

            for row in range(start, marker):
                values = rows[row]
                serial, sku, description = clean_text(values[0]), clean_text(values[1]), clean_text(values[2])
                if serial.lower() == "s/n" or clean_text(values[5]).upper() == "EU U/P":
                    continue
                if not (sku or description or re.fullmatch(r"\d+(?:\.\d+)*", serial)):
                    continue
                valid_serial = bool(re.fullmatch(r"\d+(?:\.\d+)*", serial))
                quantity = as_number(values[3])
                monetary = any(as_number(values[column]) not in (None, 0) for column in (5, 6, 7, 8))
                included = any(re.fullmatch(r"incl(?:uded|usive)?\.?", clean_text(values[column]), re.I) for column in (5, 6, 7, 8))
                transacted = monetary or quantity is not None or included
                if valid_serial:
                    is_service = product_type(sku, description) == "Professional Service"
                    if not transacted and not is_service:
                        section = sku or description
                        current = None
                        continue
                    current = new_group(row, values, serial, grouped=True)
                elif current is None or (not current["grouped"] and transacted):
                    if not transacted:
                        if sku and not description:
                            section = sku
                        continue
                    current = new_group(row, values)
                current["rows"].append(row)

            block_items: list[dict[str, Any]] = []
            for group in groups:
                cost_values: list[float | None] = []
                sell_values: list[float | None] = []
                details: list[str] = []
                for row in group["rows"]:
                    values = rows[row]
                    description = clean_text(values[2]) or clean_text(values[1])
                    if description:
                        details.append(description)
                    amounts: dict[str, float | None] = {}
                    statuses: dict[str, str] = {}
                    for side, unit_column, total_column in (("cost", 7, 9), ("sell", 6, 8)):
                        raw_total = worksheet.cell(row, total_column).value
                        raw_unit = worksheet.cell(row, unit_column).value
                        amount = as_number(raw_total)
                        status = "numeric total" if amount is not None else "missing"
                        if re.fullmatch(r"incl(?:uded|usive)?\.?", clean_text(raw_total) or clean_text(raw_unit), re.I):
                            amount, status = 0.0, "included in another price"
                        if amount is None and as_number(raw_unit) is not None and as_number(values[3]) is not None:
                            amount = as_number(raw_unit) * as_number(values[3])
                            status = "quantity × unit fallback"
                        formula = formula_sheet.cell(row, total_column).value
                        if isinstance(formula, str) and formula.startswith("=") and raw_total is None:
                            warn(row, f"{side}: formula has no cached value; recalculate and save in Excel. {status}")
                        if as_number(raw_total) is not None and as_number(raw_unit) is not None and as_number(values[3]) is not None:
                            expected = as_number(raw_unit) * as_number(values[3])
                            if abs(amount - expected) > 0.02:
                                warn(row, f"{side}: entered total {amount:.2f} differs from quantity × unit {expected:.2f}; retained entered total.")
                        amounts[side], statuses[side] = amount, status

                    kind = product_type(values[1], values[2])
                    cost_values.append(amounts["cost"])
                    sell_values.append(amounts["sell"])
                    lines.append({
                        "file": filename, "sheet": sheet_name, "block_id": block_id, "supplier": supplier,
                        "item_start_row": group["start"], "excel_row": row, "sn": clean_text(values[0]),
                        "sku": clean_text(values[1]), "description": description, "line_item_type": kind,
                        "brand_category": brand_category(values[1], values[2], kind), "quantity": as_number(values[3]),
                        "cost_raw": values[8], "sell_raw": values[7], "cost": amounts["cost"], "sell": amounts["sell"],
                        "cost_status": statuses["cost"], "sell_status": statuses["sell"],
                        "cost_currency": cost_currency, "sell_currency": sell_currency,
                    })

                cost, sell = sum_known(cost_values), sum_known(sell_values)
                kind = (category_overrides or {}).get(group["start"], product_type(group["sku"], group["name"]))
                missing_cost = sum(
                    cost_value is None and (as_number(rows[row][3]) is not None or as_number(rows[row][7]) not in (None, 0))
                    for row, cost_value in zip(group["rows"], cost_values)
                )
                missing_sell = sum(
                    sell_value is None and (as_number(rows[row][3]) is not None or as_number(rows[row][8]) not in (None, 0))
                    for row, sell_value in zip(group["rows"], sell_values)
                )
                if missing_cost or missing_sell:
                    warn(group["start"], f"Item has {missing_cost} cost and {missing_sell} sell omissions on priced/quantity rows; recorded totals may be partial.")
                record = {
                    "file": filename, "sheet": sheet_name, "customer": customer, "sales_rep": sales_rep,
                    "block_id": block_id, "supplier": supplier, "sn": group["sn"], "sku": group["sku"],
                    "item_name": group["name"], "description": "\n".join(details), "item_type": kind,
                    "brand_category": brand_category(group["sku"], group["name"], kind),
                    "segment": "Professional Services" if kind == "Professional Service" else "Hardware / Items",
                    "section": group["section"], "start_row": group["start"], "end_row": max(group["rows"]),
                    "cost_currency": cost_currency, "sell_currency": sell_currency, "cost_fx_to_sgd": cost_rate,
                    "sell_fx_to_sgd": sell_rate, "cost_fx_source": cost_fx_source, "sell_fx_source": sell_fx_source,
                    "total_cost": cost, "total_sell": sell,
                    "total_cost_sgd": cost * cost_rate if cost is not None and cost_rate is not None else None,
                    "total_sell_sgd": sell * sell_rate if sell is not None and sell_rate is not None else None,
                    "missing_cost_rows": missing_cost, "missing_sell_rows": missing_sell,
                }
                items.append(record)
                block_items.append(record)

            native_cost = sum_known(item["total_cost"] for item in block_items)
            sgd_cost = sum_known(item["total_cost_sgd"] for item in block_items)
            declared_cost = declared.get(cost_currency, declared.get("native"))
            variance = native_cost - declared_cost if native_cost is not None and declared_cost is not None else None
            if variance is not None and abs(variance) > 0.02:
                warn(marker, f"Supplier cost mismatch: extracted {native_cost:.2f}, declared {declared_cost:.2f} {cost_currency or 'native'}.")
            if declared.get("SGD") is not None and sgd_cost is not None and abs(sgd_cost - declared["SGD"]) > 0.02:
                warn(marker, "Declared SGD supplier cost does not match the converted extracted cost.")
            suppliers.append({
                "file": filename, "sheet": sheet_name, "customer": customer, "sales_rep": sales_rep,
                "block_id": block_id, "supplier": supplier, "cost_currency": cost_currency,
                "sell_currency": sell_currency, "cost_fx_to_sgd": cost_rate, "total_cost": native_cost,
                "total_sell": sum_known(item["total_sell"] for item in block_items), "total_cost_sgd": sgd_cost,
                "total_sell_sgd": sum_known(item["total_sell_sgd"] for item in block_items),
                "cost_fx_source": cost_fx_source, "sell_fx_source": sell_fx_source,
                "declared_cost": declared_cost, "cost_variance": variance,
                "missing_cost_rows": sum(item["missing_cost_rows"] for item in block_items),
                "missing_sell_rows": sum(item["missing_sell_rows"] for item in block_items),
            })
            start = block_end + 1

        for row in range(stop + 1, worksheet.max_row + 1):
            values = rows[row]
            if "subtotal" in clean_text(values[5]).lower():
                break
            if clean_text(values[2]) and any(as_number(values[column]) not in (None, 0) for column in (7, 8)):
                adjustments.append({
                    "file": filename, "sheet": sheet_name, "excel_row": row, "description": clean_text(values[2]),
                    "sell": as_number(values[7]), "cost": as_number(values[8]), "sell_currency": sell_currency,
                })
                warn(row, "Non-supplier adjustment retained separately; verify its cost currency before analysis.")

        for row in range(start, stop):
            values = rows[row]
            if (clean_text(values[1]) or clean_text(values[2])) and any(as_number(values[column]) not in (None, 0) for column in (5, 6, 7, 8)):
                warn(row, "Unassigned priced row after final supplier block; manual supplier mapping required.")

        summary_row = next(
            (row for row, values in rows.items() if "supplier price to same currency" in clean_text(values[5]).lower()),
            None,
        )
        if summary_row:
            for side, column in (("cost", 8), ("sell", 7)):
                extracted = sum_known(item[f"total_{side}_sgd"] for item in items)
                adjustment_total = sum_known(item[side] for item in adjustments)
                if side == "sell" and sell_currency == "SGD" and adjustment_total is not None:
                    extracted = (extracted or 0) + adjustment_total
                pr_summary = as_number(rows[summary_row][column])
                variance = extracted - pr_summary if extracted is not None and pr_summary is not None else None
                reconciliation.append({
                    "file": filename, "sheet": sheet_name, "side": side, "recorded_sgd": extracted,
                    "pr_summary": pr_summary, "variance": variance,
                    "missing_item_amounts": sum(item[f"total_{side}_sgd"] is None for item in items),
                })
                if variance is not None and abs(variance) > 0.02:
                    warn(summary_row, f"{side} SGD summary mismatch: extracted {extracted:.2f}, PR summary {pr_summary:.2f}.")

        return {
            "items": pd.DataFrame(items, columns=ITEM_COLUMNS),
            "suppliers": pd.DataFrame(suppliers, columns=SUPPLIER_COLUMNS),
            "lines": pd.DataFrame(lines, columns=LINE_COLUMNS),
            "issues": pd.DataFrame(issues, columns=ISSUE_COLUMNS),
            "adjustments": pd.DataFrame(adjustments, columns=ADJUSTMENT_COLUMNS),
            "reconciliation": pd.DataFrame(reconciliation, columns=RECONCILIATION_COLUMNS),
        }
    finally:
        cached.close()
        formulas.close()


def combine_results(results: Iterable[dict[str, pd.DataFrame]]) -> dict[str, pd.DataFrame]:
    """Combine multiple PR results while retaining expected columns for empty outputs."""
    results = list(results)
    definitions = {
        "items": ITEM_COLUMNS,
        "suppliers": SUPPLIER_COLUMNS,
        "lines": LINE_COLUMNS,
        "issues": ISSUE_COLUMNS,
        "adjustments": ADJUSTMENT_COLUMNS,
        "reconciliation": RECONCILIATION_COLUMNS,
    }
    combined: dict[str, pd.DataFrame] = {}
    for key, columns in definitions.items():
        frames = [result[key] for result in results if not result[key].empty]
        # Drop fields that are entirely blank in an individual source before concat.
        # They are restored below, avoiding pandas' all-NA dtype ambiguity while
        # preserving every expected output field in the combined result.
        normalized = [frame.dropna(axis=1, how="all") for frame in frames]
        combined[key] = pd.concat(normalized, ignore_index=True).reindex(columns=columns) if normalized else _empty_frame(columns)
    return combined


def error_issue(filename: str, sheet_name: str, message: str) -> pd.DataFrame:
    return pd.DataFrame([{
        "file": filename, "sheet": sheet_name, "excel_row": None, "severity": "Error", "issue": message,
    }], columns=ISSUE_COLUMNS)
